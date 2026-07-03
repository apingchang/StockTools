"""
stocktool.paper_excel - Excel 寬鬆解析
======================================

【V1.2.0-paper-trading】

設計原則：
- 只 fuzzy match「股票代號」欄位
- 其他欄位一律由 fetch_market.py 重抓（不走 Excel）
- 抓不到的（已下市/暫停交易）→ 回傳清單時跳過 + 警告
"""

from __future__ import annotations

import os
import re
from typing import List, Tuple, Optional
import openpyxl


# 股票代號的可能欄位名（含英文）
_STOCK_CODE_KEYWORDS = [
    "股票代號", "代號", "股號", "stock_code", "stock_id",
    "code", "ticker", "symbol", "證券代號"
]


def _normalize(s: str) -> str:
    return re.sub(r"\s+|_", "", str(s or "").strip().lower())


def _find_stock_code_col(headers: List[str]) -> Optional[int]:
    """Fuzzy match 哪一欄是「股票代號」"""
    for kw in _STOCK_CODE_KEYWORDS:
        kw_norm = _normalize(kw)
        for i, h in enumerate(headers):
            if _normalize(h) == kw_norm:
                return i
    # 寬鬆版：contains
    for kw in _STOCK_CODE_KEYWORDS:
        kw_norm = _normalize(kw)
        for i, h in enumerate(headers):
            if kw_norm in _normalize(h):
                return i
    return None


def _is_valid_tw_code(code: str) -> bool:
    """台股代號：4-6 位數字"""
    if not code:
        return False
    s = str(code).strip()
    if not s.isdigit():
        return False
    n = len(s)
    return 4 <= n <= 6


def parse_excel_to_stock_pool(xlsx_path: str) -> Tuple[List[str], List[str]]:
    """
    解析 Excel，回傳 (股票清單, 警告訊息清單)

    只取「股票代號」欄、其他欄位忽略（由 app 重抓）
    """
    warnings: List[str] = []
    if not os.path.exists(xlsx_path):
        return [], [f"❌ 檔案不存在: {xlsx_path}"]

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as e:
        return [], [f"❌ Excel 開啟失敗: {e}"]

    # 用第一個 sheet
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        wb.close()
        return [], ["❌ Excel 是空的"]

    headers = [str(c or "").strip() for c in rows[0]]
    code_col = _find_stock_code_col(headers)

    if code_col is None:
        wb.close()
        return [], [
            f"❌ 找不到「股票代號」欄位。表頭為：{headers}",
            f"   接受的名稱：{_STOCK_CODE_KEYWORDS}",
        ]

    seen = set()
    pool: List[str] = []
    skipped = 0

    for row in rows[1:]:
        if code_col >= len(row):
            continue
        code = row[code_col]
        if code is None:
            continue
        code_str = str(code).strip()
        if not _is_valid_tw_code(code_str):
            skipped += 1
            continue
        if code_str in seen:
            continue
        seen.add(code_str)
        pool.append(code_str)

    wb.close()

    if skipped:
        warnings.append(f"⚠️ 跳過 {skipped} 筆非台股代號格式")

    return pool, warnings


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("usage: python -m stocktool.paper_excel <xlsx_path>")
        sys.exit(1)
    pool, warns = parse_excel_to_stock_pool(sys.argv[1])
    print(f"✅ 解析出 {len(pool)} 檔：{pool}")
    for w in warns:
        print(f"  {w}")